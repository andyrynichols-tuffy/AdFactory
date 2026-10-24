"""The Scripts & Links spreadsheet (OneDrive) is the shared database: every tab is kept in
step with the app, both ways, so either of you can work in Excel (desktop, web or phone)
or in the app.

Tabs: the first one (Shoot | Link | Script, one row per video) gains Product, Account,
Creator, Tags, Due, Notes and the app's Stage / Videos / Posted; then Products, Orders,
Accounts, People, Script library, Tags and Posted videos. Every row carries an ID (last
column) so rows can be sorted, renamed and moved. Names are used for links between tabs
(type a product's name in Orders; a new name makes a new product).

How a pass works (every ~30 s, and right after a change in the app): read the workbook,
compare each cell with what it was after the last pass (AD_DATA/sheetdb.json): a cell
edited in the sheet goes into the app (the sheet wins when both changed), everything else
is written from the app. A row deleted in the sheet moves the record to the app's trash
(never when most of a tab vanished at once); a video row deleted only stops being listed.
The file is backed up before every write (AD_DATA/sheet-backups, last 30). If Excel holds
the file, the write waits for the next pass. Shoot / Link / Script of existing rows belong to
the sheet (adcut/sheet.py turns them into shoots); the app only fills them for new rows.
Never passwords: Accounts has the handle and login email only."""
import json
import re
import shutil
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path

from . import sheet, shop, store

STATE_FILE = store.AD_DATA / "sheetdb.json"
BACKUPS = store.AD_DATA / "sheet-backups"
STATUS = {"state": "starting", "msg": "", "synced": None, "path": None, "error": None, "changes": 0}
_lock = threading.Lock()

ENUMS = {
    "sample": {"free": "Free sample", "refundable": "Refundable", "paid": "Bought"},
    "status": {"ordered": "Ordered", "shipped": "Shipped", "delivered": "Delivered", "kept": "Kept",
               "returned": "Returned", "cancelled": "Cancelled"},
    "login": {"email": "Email", "Google": "Google", "Apple": "Apple", "phone": "Phone", "other": "Other"},
}
# (kind, tab, the column a new row needs, [(header, field, type, width)]); "ro" = written by the app only
TABS = [
    ("products", "Products", "name", [
        ("Product", "name", "text", 34), ("Brand", "brand", "text", 16), ("Category", "category", "text", 14),
        ("Shop link", "link", "text", 34), ("Price", "price", "text", 9), ("Commission", "commission", "text", 11),
        ("Script fill-ins", "fills", "fills", 38), ("Picture link", "image", "text", 20), ("Notes", "notes", "long", 30),
        ("Archived", "archived", "bool", 9), ("Videos made", "_made", "ro", 11), ("Posted", "_posted", "ro", 8),
        ("Views", "_views", "ro", 10)]),
    ("creators", "People", "name", [("Name", "name", "text", 20), ("Notes", "notes", "long", 40)]),
    ("accounts", "Accounts", "name", [
        ("Account", "name", "text", 22), ("Handle", "handle", "text", 18), ("Login email", "email", "text", 28),
        ("Signs in with", "login", "enum:login", 13), ("Usual creator", "creator", "ref:creators", 15),
        ("Notes", "notes", "long", 30), ("Followers", "_followers", "ro", 11)]),
    ("tags", "Tags", "name", [("Tag", "name", "text", 20), ("Group", "group", "text", 10),
                              ("What it means", "description", "long", 80)]),
    ("scripts", "Script library", "text", [
        ("Title", "title", "text", 26), ("Script", "text", "long", 70), ("Product", "product", "ref:products", 24),
        ("Tags", "tags", "tags", 28), ("Notes", "notes", "long", 26)]),
    ("orders", "Orders", "product", [
        ("Product", "product", "ref:products", 30), ("Account", "account", "ref:accounts", 18),
        ("Kind", "sample", "enum:sample", 13), ("Status", "status", "enum:status", 11), ("Order #", "order_no", "text", 16),
        ("Ordered", "ordered", "date", 13), ("Delivered", "delivered", "date", 13),
        ("Return window (days)", "return_days", "int", 11), ("Last day to return", "return_by", "date", 15),
        ("Video due", "video_due", "date", 13), ("Cost", "cost", "text", 9), ("Notes", "notes", "long", 30)]),
]
VIDEO_COLS = [("Product", "product", "ref:products", 24), ("Account", "account", "ref:accounts", 16),
              ("Creator", "creator", "ref:creators", 14), ("Tags", "tags", "tags", 26), ("Due", "due", "date", 12),
              ("Notes", "notes", "long", 26), ("Stage", "_stage", "ro", 11), ("Videos", "_videos", "ro", 8),
              ("Posted", "_posted", "ro", 8)]
POSTED_TAB = "Posted videos"
POSTED_COLS = [("Shoot", "_title", "ro", 24), ("Video", "_n", "ro", 7), ("Account", "account", "ref:accounts", 18),
               ("Product", "product", "ref:products", 26), ("TikTok link", "url", "text", 44),
               ("Posted on", "_posted_at", "ro", 12), ("Views", "_views", "ro", 9), ("Likes", "_likes", "ro", 8),
               ("Comments", "_comments", "ro", 10), ("Shares", "_shares", "ro", 8), ("Saves", "_saves", "ro", 8)]
REF_LISTS = {"products": "Products!$A$2:$A$2000", "accounts": "Accounts!$A$2:$A$500", "creators": "People!$A$2:$A$500"}


# ---------------------------------------------------------------- values

def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _date(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):  # an Excel day number
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%b %d, %Y", "%b %d %Y", "%B %d, %Y", "%B %d %Y", "%m/%d", "%b %d"):
        try:
            dt = datetime.strptime(s, f)
            if "%Y" not in f and "%y" not in f:
                dt = dt.replace(year=date.today().year)
            return dt.date().isoformat()
        except ValueError:
            pass
    return s


def _norm(typ, v) -> str:
    """A cell (or the app's value rendered as a cell) as comparable text."""
    if typ == "date":
        return _date(v)
    if typ == "bool":
        return "yes" if _s(v).lower() in ("yes", "y", "true", "x", "✓", "1") or v is True else ""
    if typ in ("long", "fills", "tags"):
        return re.sub(r"[ \t]+\n", "\n", _s(v).replace("\r\n", "\n"))
    return _s(v)


class Refs:
    """Names <-> ids for the links between tabs; unknown names make new records."""

    def __init__(self):
        self.reload()

    def reload(self):
        self.d = shop.load()

    def name(self, kind, oid):
        x = shop.by_id(self.d, kind, oid)
        return (x.get("name") or x.get("title") or "") if x else ""

    def find(self, kind, name, create=True):
        n = re.sub(r"\s+", " ", name or "").strip()
        if not n:
            return None
        key = n.lstrip("@").lower()
        for x in self.d[kind]:
            if (x.get("name") or "").strip().lower() == key or (kind == "accounts" and (x.get("handle") or "").lstrip("@").lower() == key):
                return x["id"]
        if not create:
            return None
        extra = {"description": "Added in the spreadsheet", "group": "Format"} if kind == "tags" else {}
        new = shop.upsert(kind, {"name": n, **extra})
        self.reload()
        return new["id"]


def _cell(typ, val, refs):
    """App value -> what goes in the cell."""
    if typ == "date":
        try:
            return date.fromisoformat(val) if val else None
        except ValueError:
            return val
    if typ == "int":
        return val if val not in ("", None) else None
    if typ == "bool":
        return "yes" if val else ""
    if typ.startswith("enum:"):
        return ENUMS[typ[5:]].get(val, val or "")
    if typ.startswith("ref:"):
        return refs.name(typ[4:], val)
    if typ == "tags":
        return ", ".join(n for n in (refs.name("tags", t) for t in val or []) if n)
    if typ == "fills":
        return "\n".join(f"{k}: {v}" for k, v in (val or {}).items())
    return "" if val is None else val


_SKIP = object()


def _to_app(typ, v, refs):
    """Cell -> app value (_SKIP = can't read it, leave the app's value alone)."""
    s = _norm(typ, v)
    if typ == "date":
        return s if (not s or re.fullmatch(r"\d{4}-\d\d-\d\d", s)) else _SKIP
    if typ == "int":
        return int(float(s)) if re.fullmatch(r"\d+(\.0+)?", s) else (None if not s else _SKIP)
    if typ == "bool":
        return s == "yes"
    if typ.startswith("enum:"):
        rev = {lab.lower(): k for k, lab in ENUMS[typ[5:]].items()} | {k.lower(): k for k in ENUMS[typ[5:]]}
        return rev.get(s.lower(), _SKIP) if s else _SKIP
    if typ.startswith("ref:"):
        return refs.find(typ[4:], s) if s else None
    if typ == "tags":
        return [i for i in (refs.find("tags", n) for n in re.split(r"[,;\n]", s)) if i]
    if typ == "fills":
        out = {}
        for part in re.split(r"[;\n]", s):
            if ":" in part:
                k, val = part.split(":", 1)
                if k.strip() and val.strip():
                    out[k.strip().lower()] = val.strip()
        return out
    return s


# ---------------------------------------------------------------- the workbook

def _table(ws):
    heads = {}
    for c in ws[1]:
        if _s(c.value):
            heads.setdefault(_s(c.value).lower(), c.column)
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = {h: ws.cell(r, ci).value for h, ci in heads.items()}
        if any(_s(v) for v in vals.values()):
            rows.append((r, vals))
    return heads, rows


def _headers(ws, cols, new_tab=False):
    """Make sure every column (and ID, last) has a header; returns {header lower: column}."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    heads, _ = _table(ws)
    last = max([0] + list(heads.values()))
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E9ECEF")
    for h, _, typ, width in cols + [("ID", "id", "ro", 9)]:
        if h.lower() in heads:
            continue
        last += 1
        c = ws.cell(1, last, h)
        c.font, c.fill = bold, fill
        if h == "ID":
            c.font = Font(bold=True, color="868E96")
        L = get_column_letter(last)
        ws.column_dimensions[L].width = width
        lst = None
        if typ.startswith("enum:"):
            lst = '"' + ",".join(ENUMS[typ[5:]].values()) + '"'
        elif typ == "bool":
            lst = '"yes"'
        elif typ.startswith("ref:"):
            lst = "=" + REF_LISTS[typ[4:]]
        if lst:  # a dropdown, but any text is fine (a new name makes a new record)
            dv = DataValidation(type="list", formula1=lst, allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(f"{L}2:{L}2000")
        heads[h.lower()] = last
    if new_tab:
        ws.freeze_panes = "A2"
    return heads


class Book:
    def __init__(self, wb):
        self.wb = wb
        self.dirty = 0

    def put(self, ws, r, col, typ, value):
        from openpyxl.styles import Alignment
        c = ws.cell(r, col)
        if _norm(typ, c.value) == _norm(typ, value):
            return
        c.value = value if value != "" else None
        if typ == "date" and isinstance(value, date):
            c.number_format = "mmm d, yyyy"
        if typ in ("long", "fills"):
            c.alignment = Alignment(wrap_text=True, vertical="top")
        self.dirty += 1

    def tab(self, title):
        if title in self.wb.sheetnames:
            return self.wb[title], False
        self.dirty += 1
        return self.wb.create_sheet(title), True


# ---------------------------------------------------------------- merging

def _merge_kind(book, refs, base, kind, title, need, cols, extra):
    """One app collection <-> one tab."""
    ws, new = book.tab(title)
    heads = _headers(ws, cols, new_tab=new)
    if new:
        book.dirty += 1
    _, rows = _table(ws)
    b = base.setdefault(title, {})
    idc = heads["id"]
    listed = lambda x: not (kind == "scripts" and x.get("from_shoot"))  # those are the first tab's rows
    app = {x["id"]: x for x in refs.d[kind] if listed(x)}
    seen, row_of = set(), {}
    for r, vals in rows:
        rid = _s(vals.get("id"))
        if rid in app:
            seen.add(rid)
            row_of[rid] = r
            rec, patch = app[rid], {}
            for h, key, typ, _ in cols:
                if typ == "ro":
                    continue
                sv = _norm(typ, vals.get(h.lower()))
                av = _norm(typ, _cell(typ, rec.get(key), refs))
                bv = b.get(rid, {}).get(h)
                if sv != av and (sv != bv if bv is not None else sv != ""):  # never synced: only a filled-in cell counts
                    v = _to_app(typ, vals.get(h.lower()), refs)
                    if v is not _SKIP:
                        patch[key] = v
            if patch:
                if kind == "scripts" and "text" in patch:
                    patch["edited"] = True
                shop.upsert(kind, {"id": rid, **patch})
                STATUS["changes"] += 1
        elif rid and rid in b:
            row_of[rid] = r  # removed in the app: the row goes too (below)
        else:  # a row typed into the sheet
            needed = next((h for h, key, *_ in cols if key == need), None)
            if not _s(vals.get(needed.lower())):
                continue
            rec = {}
            for h, key, typ, _ in cols:
                if typ != "ro" and _s(vals.get(h.lower())):
                    v = _to_app(typ, vals.get(h.lower()), refs)
                    if v is not _SKIP:
                        rec[key] = v
            if kind == "tags" and not rec.get("description"):
                rec["description"] = "Added in the spreadsheet"
            made = shop.upsert(kind, rec)
            refs.reload()
            ws.cell(r, idc, made["id"])
            book.dirty += 1
            seen.add(made["id"])
            row_of[made["id"]] = r
            STATUS["changes"] += 1
    gone = [i for i in b if i in app and i not in seen]
    if gone and len(gone) <= max(2, len(b) // 2):  # deleted in the sheet -> the app's trash
        for i in gone:
            shop.remove(kind, i)
            STATUS["changes"] += 1
        refs.reload()
    # app -> sheet
    app = {x["id"]: x for x in refs.d[kind] if listed(x)}
    nb, last = {}, ws.max_row
    for oid, rec in app.items():
        r = row_of.get(oid)
        if r is None:
            last += 1
            r = last
            book.put(ws, r, idc, "text", oid)
        snap = {}
        for h, key, typ, _ in cols:
            val = extra(rec, key) if typ == "ro" else _cell(typ, rec.get(key), refs)
            book.put(ws, r, heads[h.lower()], typ, val)
            snap[h] = _norm(typ, val)
        nb[oid] = snap
    for oid in sorted((i for i in row_of if i not in app), key=lambda i: -row_of[i]):
        ws.delete_rows(row_of[oid])
        book.dirty += 1
    base[title] = nb


def _merge_videos(book, refs, base):
    """The first tab: one row per video plan (shoot)."""
    ws = book.wb.worksheets[0]
    heads = _headers(ws, VIDEO_COLS)
    _, rows = _table(ws)
    b = base.setdefault("__videos__", {})
    shoots = {sid: s for sid, s, _ in shop._all_shoots()}
    board = {r["sid"]: r for r in shop.board()}
    key = lambda n: re.sub(r"\s+", " ", _s(n)).lower()
    by_row = {key((s.get("sheet") or {}).get("row")): sid for sid, s in shoots.items() if (s.get("sheet") or {}).get("row")}
    seen, row_of = set(), {}
    for r, vals in rows:
        rid = _s(vals.get("id"))
        sid = rid if rid in shoots else by_row.get(key(vals.get("shoot")))
        if not sid or sid in seen:
            continue  # adcut/sheet.py makes the shoot for a new row; next pass links it
        seen.add(sid)
        row_of[sid] = r
        if rid != sid:
            book.put(ws, r, heads["id"], "text", sid)
        plan, patch = shop.plan_of(shoots[sid]), {}
        for h, k, typ, _ in VIDEO_COLS:
            if typ == "ro":
                continue
            sv = _norm(typ, vals.get(h.lower()))
            av = _norm(typ, _cell(typ, plan.get(k), refs))
            bv = b.get(sid, {}).get(h)
            if sv != av and (sv != bv if bv is not None else sv != ""):  # never synced: only a filled-in cell counts
                v = _to_app(typ, vals.get(h.lower()), refs)
                if v is not _SKIP:
                    patch[k] = v
        if patch:
            shop.set_plan(sid, patch)
            shoots[sid] = store.load(sid)
            STATUS["changes"] += 1
    last = ws.max_row
    for sid, s in shoots.items():
        if sid in seen or not s.get("shop") or (s.get("sheet") or {}).get("hidden"):
            continue
        if sid in b:  # its row was deleted in the sheet: stop listing it (the shoot stays)
            store.update(sid, lambda x: x.setdefault("sheet", {}).update(hidden=True))
            continue
        plan = shop.plan_of(s)
        sc = shop.by_id(refs.d, "scripts", plan["script"])
        text = sc["text"] if sc else ""
        last += 1
        r = last
        link = s["reference"].get("url") or ""
        book.put(ws, r, heads["shoot"], "text", s["title"])
        book.put(ws, r, heads["link"], "text", link)
        book.put(ws, r, heads["script"], "long", text)
        book.put(ws, r, heads["id"], "text", sid)
        row_of[sid] = r
        lines = [l["text"] for l in s["script"]]

        def mark(x, row={"link": link, "script": text.strip()}, title=s["title"], lines=lines):
            m = x.setdefault("sheet", {})  # so adcut/sheet.py sees this row as already in sync
            m.update(row=title, sig=sheet._sig(row), lines=lines)
        store.update(sid, mark)
        shoots[sid] = store.load(sid)
    stage = dict(shop.STAGES)
    nb = {}
    for sid, r in row_of.items():
        plan, bo, snap = shop.plan_of(shoots[sid]), board.get(sid, {}), {}
        for h, k, typ, _ in VIDEO_COLS:
            val = {"_stage": stage.get(bo.get("stage"), ""), "_videos": bo.get("videos", 0),
                   "_posted": bo.get("posted", 0)}.get(k) if typ == "ro" else _cell(typ, plan.get(k), refs)
            book.put(ws, r, heads[h.lower()], typ, val)
            snap[h] = _norm(typ, val)
        nb[sid] = snap
    base["__videos__"] = nb


def _merge_posted(book, refs, base):
    ws, new = book.tab(POSTED_TAB)
    heads = _headers(ws, POSTED_COLS, new_tab=new)
    _, rows = _table(ws)
    b = base.setdefault(POSTED_TAB, {})
    posts = {f"{p['sid']}#{p['n']}": p for p in shop.postings() if p["sid"] and (p["posted"] or p["url"])}
    row_of = {}
    for r, vals in rows:
        rid = _s(vals.get("id"))
        if rid not in posts:
            continue
        row_of[rid] = r
        p, patch = posts[rid], {}
        for h, k, typ, _ in POSTED_COLS:
            if typ == "ro":
                continue
            sv = _norm(typ, vals.get(h.lower()))
            av = _norm(typ, _cell(typ, p.get(k), refs))
            bv = b.get(rid, {}).get(h)
            if sv != av and (sv != bv if bv is not None else sv != ""):  # never synced: only a filled-in cell counts
                v = _to_app(typ, vals.get(h.lower()), refs)
                if v is not _SKIP:
                    patch[k] = v
        if patch:
            shop.posting(p["sid"], {"n": p["n"], **patch})
            STATUS["changes"] += 1
    posts = {f"{p['sid']}#{p['n']}": p for p in shop.postings() if p["sid"] and (p["posted"] or p["url"])}
    nb, last = {}, ws.max_row
    for rid, p in sorted(posts.items(), key=lambda kv: kv[1].get("posted_at") or 0):
        r = row_of.get(rid)
        if r is None:
            last += 1
            r = last
            book.put(ws, r, heads["id"], "text", rid)
        st = p.get("stats") or {}
        ro = {"_title": p["title"], "_n": p["n"], "_posted_at": date.fromtimestamp(p["posted_at"]) if p.get("posted_at") else None,
              **{f"_{k}": st.get(k) for k in ("views", "likes", "comments", "shares", "saves")}}
        snap = {}
        for h, k, typ, _ in POSTED_COLS:
            val = ro.get(k) if typ == "ro" else _cell(typ, p.get(k), refs)
            book.put(ws, r, heads[h.lower()], "date" if k == "_posted_at" else typ, val)
            snap[h] = _norm(typ, val)
        nb[rid] = snap
    base[POSTED_TAB] = nb


# ---------------------------------------------------------------- one pass

def _backup(path):
    BACKUPS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, BACKUPS / f"{time.strftime('%Y%m%d-%H%M%S')} {path.name}")
    for old in sorted(BACKUPS.glob("*.xlsx"))[:-30]:
        old.unlink()


def sync(path=None) -> dict:
    """One pass both ways. Safe to call often; writes only when something differs."""
    import openpyxl
    if not _lock.acquire(blocking=False):
        return STATUS
    try:
        path = Path(path) if path else sheet.find()
        STATUS["path"] = str(path) if path else None
        if not path:
            STATUS.update(state="missing", msg="No “Scripts & Links” spreadsheet found in OneDrive")
            return STATUS
        try:
            wb = openpyxl.load_workbook(path)
        except Exception as e:  # syncing / open elsewhere: next pass
            STATUS.update(state="waiting", msg=f"Couldn't open the spreadsheet yet ({type(e).__name__})", error=str(e)[:200])
            return STATUS
        try:
            state = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            state = {}
        base = state.setdefault("base", {}) if state.get("path") in (None, str(path)) else {}
        state["path"] = str(path)
        book, refs = Book(wb), Refs()
        summary = shop.summary()
        stats = shop.load_stats()

        def extra(rec, k):
            if k in ("_made", "_posted", "_views"):
                per = summary["per"]["products"].get(rec["id"]) or {}
                return {"_made": (summary["made"].get(rec["id"]) or {}).get("videos", 0), "_posted": per.get("posted", 0),
                        "_views": per.get("views") or None}[k]
            if k == "_followers":
                return shop.account_numbers(stats, rec).get("followers")
            return None
        for kind, title, need, cols in TABS:
            _merge_kind(book, refs, base, kind, title, need, cols, extra)
        _merge_videos(book, refs, base)
        _merge_posted(book, refs, base)
        if book.dirty:
            tmp = store.AD_DATA / "sheetdb-write.xlsx"
            wb.save(tmp)
            try:
                _backup(path)
                shutil.copyfile(tmp, path)  # in place, so OneDrive keeps it as the same file
            except PermissionError:
                STATUS.update(state="waiting", msg="The spreadsheet is open in Excel on this PC: the app's changes go in once it's closed (or saved)")
                return STATUS
        state["base"] = base
        STATE_FILE.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
        STATUS.update(state="ok", msg="", error=None, synced=time.time(), written=bool(book.dirty))
        return STATUS
    except Exception as e:  # shown in the app
        STATUS.update(state="error", msg=f"Spreadsheet sync failed: {e}", error=str(e)[:300])
        return STATUS
    finally:
        _lock.release()
